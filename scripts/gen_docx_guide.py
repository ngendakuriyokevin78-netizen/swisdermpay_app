from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "CashTel_Guide_Utilisation_et_Tests.docx"

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def title(t):
    p = doc.add_heading(t, level=0)
    return p

def h1(t):
    doc.add_heading(t, level=1)

def h2(t):
    doc.add_heading(t, level=2)

def para(t, bold=False, italic=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    return p

def bullets(items):
    for it in items:
        doc.add_paragraph(it, style='List Bullet')

def code(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
    pf = p.paragraph_format
    pf.left_indent = Inches(0.2)
    return p

def flow(steps):
    n = len(steps)
    cols = n * 2 - 1
    tbl = doc.add_table(rows=1, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = 'Table Grid'
    for i, s in enumerate(steps):
        c = i * 2
        cell = tbl.cell(0, c)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(s)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(cell, "1F4E79")
        if i < n - 1:
            ac = tbl.cell(0, c + 1)
            ac.text = ""
            ap = ac.paragraphs[0]
            ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ar = ap.add_run("  \u25B6  ")
            ar.bold = True
            ar.font.size = Pt(12)
            ar.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            shade(ac, "D9E2F3")
    doc.add_paragraph("")

def endpoint_table(rows):
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = 'Table Grid'
    hdr = tbl.rows[0].cells
    hdr[0].text = 'Methode'
    hdr[1].text = 'URL'
    hdr[2].text = 'Body / Auth'
    for c in hdr:
        shade(c, "1F4E79")
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for m, u, b in rows:
        cells = tbl.add_row().cells
        cells[0].text = m
        cells[1].text = u
        cells[2].text = b
    doc.add_paragraph("")

title("Cash Tel API — Guide d'utilisation et de tests")
p = doc.add_paragraph()
r = p.add_run("Mobile Money Burundi (BIF) — Django 5 + DRF + JWT + Swagger\nBase URL : http://127.0.0.1:8000 — Docs : /api/docs/")
r.italic = True
para("Objectif : utiliser l'API etape par etape et savoir tester chaque module (Swagger, curl, Postman, pytest).", italic=True)

h1("1. Demarrage rapide")
bullets([
    "Copier .env : cp .env.example .env",
    "Migrer : python manage.py migrate",
    "Donnees initiales : python manage.py create_initial_data",
    "Lancer : python manage.py runserver",
    "Docker : docker-compose up --build -> API http://localhost:8000/api/docs/",
])
code("cp .env.example .env\npython manage.py migrate\npython manage.py create_initial_data\npython manage.py runserver")

h1("2. Vue globale — schema fleche")
para("Parcours obligatoire d'un utilisateur :", bold=True)
flow(["1. Register", "2. Verify OTP", "3. Login (JWT)", "4. Wallet / Transfer / Bills", "5. Historique"])
para("Authentification : toutes les routes sauf register / verify-otp / resend-otp / login exigent :", bold=True)
code("Authorization: Bearer <access_token>\nContent-Type: application/json")

h1("3. Authentification etape par etape")
h2("3.1 Inscription -> OTP -> Login")
flow(["POST /register", "POST /verify-otp", "POST /login", "Bearer JWT"])
bullets([
    "POST /api/auth/register/ : {phone_number, first_name, last_name, pin(4 chiffres)} -> 201 + OTP envoye (logs en mode mock).",
    "POST /api/auth/verify-otp/ : {phone_number, otp(6 chiffres)} -> telephone verifie.",
    "POST /api/auth/resend-otp/ : {phone_number} si OTP expire.",
    "POST /api/auth/login/ : {phone_number, pin} -> {access, refresh, user}. Erreurs : 401 PIN faux, 403 bloque (3 echecs) / non verifie.",
])
code('curl -X POST http://127.0.0.1:8000/api/auth/register/ -H "Content-Type: application/json" -d \'{"phone_number":"+25762000001","first_name":"Alice","last_name":"Test","pin":"1234"}\'')
code('curl -X POST http://127.0.0.1:8000/api/auth/verify-otp/ -H "Content-Type: application/json" -d \'{"phone_number":"+25762000001","otp":"123456"}\'')
code('curl -X POST http://127.0.0.1:8000/api/auth/login/ -H "Content-Type: application/json" -d \'{"phone_number":"+25762000001","pin":"1234"}\'')

h2("3.2 Profil, PIN, Session")
flow(["Login", "GET /profile", "POST /change-pin", "POST /logout"])
bullets([
    "GET /api/auth/profile/ (Bearer) -> profil + wallet_id + balance.",
    "POST /api/auth/change-pin/ : {current_pin, new_pin, confirm_pin}.",
    "POST /api/auth/token/refresh/ : {refresh} -> nouveau access.",
    "POST /api/auth/logout/ : {refresh} -> blacklist.",
])
endpoint_table([
    ("GET", "/api/auth/profile/", "Bearer"),
    ("POST", "/api/auth/change-pin/", "{current_pin,new_pin,confirm_pin}"),
    ("POST", "/api/auth/token/refresh/", "{refresh}"),
    ("POST", "/api/auth/logout/", "{refresh} + Bearer"),
])

h1("4. Wallet")
flow(["Login", "GET /wallet/me", "GET /wallet/qr-data", "Partager qr_data"])
bullets([
    "GET /api/wallet/me/ -> {wallet_id, balance, status, qr_data}.",
    "GET /api/wallet/qr-data/ -> {qr_data: CASHTEL:<wallet_id>:<phone>} a donner a l'expediteur.",
    "GET /api/wallet/qr/ -> image PNG.",
])
code('curl http://127.0.0.1:8000/api/wallet/me/ -H "Authorization: Bearer <ACCESS>"')
code('curl http://127.0.0.1:8000/api/wallet/qr-data/ -H "Authorization: Bearer <ACCESS>"')

h1("5. Transferts P2P")
h2("5.1 Par numero")
flow(["Solde OK ?", "POST /transfer/send", "Transaction SUCCESS", "Historique"])
code('curl -X POST http://127.0.0.1:8000/api/transfer/send/ -H "Authorization: Bearer <ACCESS_A>" -H "Content-Type: application/json" -d \'{"receiver_phone":"+25762000002","amount":"5000.00","pin":"1234"}\'')
h2("5.2 Par QR")
flow(["B : GET qr-data", "A : POST /transfer/qr", "SUCCESS", "Verif soldes"])
code('curl -X POST http://127.0.0.1:8000/api/transfer/qr/ -H "Authorization: Bearer <ACCESS_A>" -H "Content-Type: application/json" -d \'{"qr_data":"CASHTEL:<wallet_id>:+25762000002","amount":"5000.00","pin":"1234"}\'')
h2("5.3 Frais")
code('curl http://127.0.0.1:8000/api/transfer/fees/ -H "Authorization: Bearer <ACCESS>"\ncurl "http://127.0.0.1:8000/api/transfer/simulate/?amount=5000" -H "Authorization: Bearer <ACCESS>"')

h1("6. Agent (Cash In / Out)")
flow(["Agent login", "POST /agent/cashin OU cashout", "Transaction", "Recu"])
bullets([
    "POST /api/agent/cashin/ (role AGENT) : {customer_phone, amount} — pas de PIN.",
    "POST /api/agent/cashout/ : {customer_phone, amount, pin}.",
])
code('curl -X POST http://127.0.0.1:8000/api/agent/cashin/ -H "Authorization: Bearer <ACCESS_AGENT>" -H "Content-Type: application/json" -d \'{"customer_phone":"+25762000001","amount":"10000.00"}\'')

h1("7. Factures (Bills)")
flow(["GET /billers", "POST /bills/pay", "Transaction BILL", "Historique"])
code('curl http://127.0.0.1:8000/api/bills/billers/ -H "Authorization: Bearer <ACCESS>"')
code('curl -X POST http://127.0.0.1:8000/api/bills/pay/ -H "Authorization: Bearer <ACCESS>" -H "Content-Type: application/json" -d \'{"biller_code":"REGIDESO","reference_number":"123456","amount":"5000.00","pin":"1234"}\'')

h1("8. Historique")
code('curl http://127.0.0.1:8000/api/transfer/history/ -H "Authorization: Bearer <ACCESS>"\ncurl "http://127.0.0.1:8000/api/transactions/?type=TRANSFER&status=SUCCESS" -H "Authorization: Bearer <ACCESS>"')
para("Filtres : ?type=TRANSFER|DEPOSIT|WITHDRAWAL|BILL &status=PENDING|SUCCESS|FAILED. Pagination 20.", italic=True)

h1("9. Comment tester")
h2("9.1 Swagger (le plus simple)")
bullets([
    "Ouvrir http://127.0.0.1:8000/api/docs/",
    "1) POST /api/auth/register/ -> Execute, 2) POST verify-otp, 3) POST login -> copier access.",
    "4) Bouton Authorize (en haut) -> coller : Bearer <access> -> tester wallet / transfer / bills.",
    "Schema fleche : Swagger -> Try it out -> Execute -> verifier code 200/201 -> Response body.",
])
h2("9.2 curl (voir sections 3-8)")
flow(["Terminal 1: runserver", "Terminal 2: curl register", "curl verify + login", "curl metier + verif"])
h2("9.3 Postman")
bullets([
    "Creer collection CashTel + variable {{base}}=http://127.0.0.1:8000, {{access}}.",
    "Requete login -> onglet Tests : pm.environment.set('access', pm.response.json().access).",
    "Autres requetes : onglet Auth -> Bearer Token -> {{access}}.",
    "Tester cas nominaux + echecs : mauvais PIN (401), compte bloque (403), montant insuffisant (400).",
])
h2("9.4 Tests automatiques")
code("pytest -v\npytest tests/test_transfer.py -v\npython scripts/test_qr_transfer.py --demo")
bullets([
    "pytest -v : lance toute la suite (6 tests transfert).",
    "scripts/test_qr_transfer.py --demo : scenario QR bout en bout.",
    "Ajouter un test : creer tests/test_<module>.py avec APIClient + JWT.",
])
h2("9.5 Checklist de test par module")
endpoint_table([
    ("Auth", "register/verify/login/profile/change-pin", "201/200, 400, 401, 403"),
    ("Wallet", "me/qr-data/qr", "200 + balance correcte"),
    ("Transfer", "send/qr/simulate/fees", "200, solde debite/credite + frais"),
    ("Agent", "cashin/cashout", "200, role AGENT requis"),
    ("Bills", "billers/pay", "200, biller_code valide"),
    ("History", "history/transactions", "200, filtres + pagination"),
])

h1("10. Erreurs et depannage")
bullets([
    "Format standard : {success:false, error:'...', details:{...}}.",
    "400 : validation (PIN non numerique, montant <=0, qr_data sans prefixe CASHTEL:).",
    "401 : JWT absent/expire ou PIN incorrect (tentatives restantes indiquees).",
    "403 : compte bloque / non verifie / desactive, ou role AGENT manquant.",
    "404 : utilisateur/wallet/biller introuvable.",
    "OTP non recu : backend mock -> lire logs runserver.",
    "401 partout : refresh le token via /token/refresh/.",
])

para("Document genere automatiquement — Cash Tel.", italic=True)
doc.save(OUT)
print(f"saved {OUT}")
